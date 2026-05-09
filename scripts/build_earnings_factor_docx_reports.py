from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(31, 78, 121)
LIGHT = "D9EAF7"


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _fmt(value, pct: bool = False, digits: int = 2) -> str:
    if value is None:
        return "n/a"
    try:
        if pd.isna(value):
            return "n/a"
        value = float(value)
    except (TypeError, ValueError):
        return str(value)
    if pct:
        return f"{value:.{digits}%}"
    return f"{value:.{digits}f}"


def _set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _style_doc(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Calibri"
        styles[style_name].font.color.rgb = ACCENT
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)


def _add_kv_table(doc: Document, rows: list[tuple[str, str]], widths: tuple[float, float] = (2.4, 4.4)) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Value"
    for cell in hdr:
        _set_cell_shading(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        row.cells[0].width = Inches(widths[0])
        row.cells[1].width = Inches(widths[1])


def _add_dataframe_table(doc: Document, frame: pd.DataFrame, max_rows: int = 12) -> None:
    if frame.empty:
        doc.add_paragraph("No rows.")
        return
    frame = frame.head(max_rows).copy()
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for idx, col in enumerate(frame.columns):
        table.rows[0].cells[idx].text = str(col)
        _set_cell_shading(table.rows[0].cells[idx], LIGHT)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = _fmt(value, pct=False, digits=4) if isinstance(value, float) else str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _signal_ic(metrics: dict, bucket: str = "selected") -> float | None:
    return (
        metrics.get("signal_diagnostics", {})
        .get(bucket, {})
        .get("rank_ic_rank_score_vs_strategy_excess_vs_QQQ_event_return")
    )


def _interpretation(manifest: dict, metrics: dict) -> str:
    name = manifest.get("factor_name", "")
    net = metrics.get("strategy_net", {})
    excess = metrics.get("excess_vs_QQQ", {})
    selected = metrics.get("event_counts", {}).get("selected_events", 0)
    sharpe = net.get("sharpe", net.get("information_ratio"))
    excess_ir = excess.get("sharpe", excess.get("information_ratio"))
    if "PEAD" in name:
        return (
            "This is the strongest current candidate in the suite. The result supports the revised research "
            "question: confirmed post-earnings drift has more signal than pre-earnings jump prediction. "
            f"The sample has {selected} selected events, positive QQQ-relative excess, and usable but still "
            "not production-grade rank IC."
        )
    if "EAP" in name:
        return (
            "This behaves as a weak baseline. The diversified announcement-premium idea is not enough by "
            "itself in the current technology sample; it needs stronger conditioning or a broader PIT universe."
        )
    if "leader_confirmed" in name:
        return (
            "This is a weak but economically coherent basket-timing candidate. QQQ-relative excess is better "
            "than the absolute return headline, which means the basket timing is partially hedging market beta, "
            "but the absolute return and Sharpe remain too low."
        )
    if "leader_follower" in name:
        return (
            "This version produces positive absolute return but weak score IC and material drawdown. It is a "
            "candidate for redesign, not a final factor: the leader-follower gap mechanism needs better timing "
            "and broader data before it can be trusted."
        )
    return (
        f"The run selected {selected} events with Sharpe {_fmt(sharpe)} and QQQ-excess IR {_fmt(excess_ir)}. "
        "Use this as a research artifact, not a production approval."
    )


def _signal_diagnostics_frame(metrics: dict) -> pd.DataFrame:
    rows = []
    diagnostics = metrics.get("signal_diagnostics", {})
    for bucket in ["eligible", "tradable", "selected"]:
        item = diagnostics.get(bucket, {})
        if not item:
            continue
        rows.append(
            {
                "bucket": bucket,
                "n": item.get("n"),
                "rank_ic_score_vs_ret": _fmt(item.get("rank_ic_rank_score_vs_strategy_event_return"), digits=3),
                "rank_ic_score_vs_excess": _fmt(
                    item.get("rank_ic_rank_score_vs_strategy_excess_vs_QQQ_event_return"), digits=3
                ),
                "rank_ic_pct_vs_excess": _fmt(
                    item.get("rank_ic_score_percentile_history_vs_strategy_excess_vs_QQQ_event_return"), digits=3
                ),
            }
        )
    return pd.DataFrame(rows)


def _add_plot(doc: Document, run_dir: Path, plot_name: str, title: str) -> None:
    path = run_dir / "plots" / plot_name
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.4))
    caption = doc.add_paragraph(title)
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True


def build_report(run_dir: Path, out_path: Path | None = None) -> Path:
    manifest = _load_json(run_dir / "manifest.json")
    metrics = _load_json(run_dir / "metrics_summary.json")
    spec = _load_json(run_dir / "experiment_spec.json")
    diagnostics = _load_json(run_dir / "data_diagnostics.json")
    yearly = pd.read_csv(run_dir / "metrics_by_year.csv") if (run_dir / "metrics_by_year.csv").exists() else pd.DataFrame()
    selected = pd.read_csv(run_dir / "selected_events.csv") if (run_dir / "selected_events.csv").exists() else pd.DataFrame()

    doc = Document()
    _style_doc(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(manifest["factor_name"])
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Qlib-backed Earnings Factor Research Report").italic = True

    doc.add_heading("Research Decision", level=1)
    selected_events = metrics.get("event_counts", {}).get("selected_events")
    net = metrics.get("strategy_net", {})
    excess = metrics.get("excess_vs_QQQ", {})
    selected_event = metrics.get("selected_event", {})
    selected_ic = _signal_ic(metrics, "selected")
    decision_rows = [
        ("Research status", manifest.get("status", "research_only")),
        ("Selected events", selected_events),
        ("Net cumulative return", _fmt(net.get("cumulative_return"), pct=True)),
        ("Net annualized return", _fmt(net.get("annualized_return"), pct=True)),
        ("Net Sharpe / IR", _fmt(net.get("sharpe", net.get("information_ratio")))),
        ("Max drawdown", _fmt(net.get("max_drawdown"), pct=True)),
        ("Exposure-matched excess vs QQQ", _fmt(excess.get("cumulative_return"), pct=True)),
        ("Excess Sharpe / IR", _fmt(excess.get("sharpe", excess.get("information_ratio")))),
        ("Avg event return after cost", _fmt(selected_event.get("average_event_return_after_cost"), pct=True)),
        ("Avg event excess vs QQQ", _fmt(selected_event.get("average_excess_vs_QQQ_event_return"), pct=True)),
        ("Selected rank IC vs QQQ excess", _fmt(selected_ic, digits=3)),
    ]
    _add_kv_table(doc, [(str(k), str(v)) for k, v in decision_rows])

    doc.add_heading("Research Interpretation", level=1)
    doc.add_paragraph(_interpretation(manifest, metrics))

    doc.add_heading("Economic Hypothesis", level=1)
    doc.add_paragraph(spec.get("hypothesis", "n/a"))
    doc.add_heading("Factor Expression", level=1)
    for key, value in spec.get("formula", {}).items():
        p = doc.add_paragraph()
        p.add_run(f"{key}: ").bold = True
        p.add_run(str(value))

    doc.add_heading("Experiment Design", level=1)
    execution = spec.get("execution", {})
    universe = spec.get("universe", {})
    _add_kv_table(
        doc,
        [
            ("Date range", f"{spec.get('date_range', {}).get('start')} to {spec.get('date_range', {}).get('end')}"),
            ("Benchmark", str(spec.get("benchmark"))),
            ("Universe size", str(len(universe.get("symbols", [])))),
            ("Entry", str(execution.get("entry"))),
            ("Exit", str(execution.get("exit"))),
            ("Hold days", str(execution.get("hold_days"))),
            ("Max active positions", str(execution.get("max_active_positions"))),
            ("Max weight", str(execution.get("max_weight"))),
            ("One-way cost bps", str(spec.get("cost_model", {}).get("one_way_cost_bps"))),
        ],
    )

    doc.add_heading("Signal Diagnostics / IC", level=1)
    doc.add_paragraph(
        "Rank IC is computed against the actual strategy holding-horizon return, not just the earnings "
        "announcement window. This makes PEAD and basket variants comparable despite different horizons."
    )
    _add_dataframe_table(doc, _signal_diagnostics_frame(metrics), max_rows=6)

    doc.add_heading("Performance Charts", level=1)
    for plot_name, caption in [
        ("cumulative_return.png", "Cumulative return: strategy, exposure-matched QQQ, and excess return."),
        ("drawdown.png", "Net strategy drawdown."),
        ("yearly_returns.png", "Yearly strategy and excess returns."),
        ("exposure.png", "Daily portfolio exposure."),
        ("event_return_distribution.png", "Selected event return distribution after cost."),
        ("selected_trade_timeline.png", "Selected trade timeline."),
        ("score_vs_event_return.png", "Score versus realized event return."),
        ("top_bottom_trades.png", "Top and bottom selected trades."),
    ]:
        _add_plot(doc, run_dir, plot_name, caption)

    doc.add_heading("Metrics by Year", level=1)
    if not yearly.empty:
        cols = [
            col
            for col in [
                "year",
                "strategy_net_return",
                "excess_vs_QQQ_year_return",
                "max_drawdown",
                "sharpe",
                "avg_exposure",
                "selected_events",
                "avg_event_return_after_cost",
            ]
            if col in yearly.columns
        ]
        table_yearly = yearly[cols].copy()
        for col in table_yearly.columns:
            if col != "year" and col != "selected_events":
                table_yearly[col] = table_yearly[col].map(lambda value: _fmt(value, pct=True) if "return" in col or "drawdown" in col or "exposure" in col else _fmt(value))
        _add_dataframe_table(doc, table_yearly, max_rows=20)
    else:
        doc.add_paragraph("No yearly metrics.")

    doc.add_heading("Selected Trades", level=1)
    if not selected.empty:
        for col in ["event_return_after_cost", "excess_vs_QQQ_event_return", "rank_score"]:
            if col in selected:
                selected[col] = pd.to_numeric(selected[col], errors="coerce")
        top = selected.sort_values("event_return_after_cost", ascending=False).head(8)
        bottom = selected.sort_values("event_return_after_cost", ascending=True).head(8)
        trade_cols = [col for col in ["symbol", "entry_date", "exit_date", "event_return_after_cost", "excess_vs_QQQ_event_return", "rank_score"] if col in selected.columns]
        doc.add_heading("Top trades", level=2)
        _add_dataframe_table(doc, top[trade_cols], max_rows=8)
        doc.add_heading("Bottom trades", level=2)
        _add_dataframe_table(doc, bottom[trade_cols], max_rows=8)
    else:
        doc.add_paragraph("No selected trades.")

    doc.add_heading("Literature Support", level=1)
    for item in spec.get("literature", []):
        p = doc.add_paragraph(style=None)
        p.add_run(item.get("title", "")).bold = True
        p.add_run(f" - {item.get('authors', '')}. ")
        p.add_run(item.get("url", ""))

    doc.add_heading("Data and Research Caveats", level=1)
    _add_kv_table(
        doc,
        [
            ("Requested symbols", str(diagnostics.get("requested_symbol_count"))),
            ("Available price symbols", str(diagnostics.get("available_price_symbol_count"))),
            ("Event count", str(diagnostics.get("event_count"))),
            ("Missing price symbols", ", ".join(diagnostics.get("missing_price_symbols", [])) or "none"),
            ("Missing earnings symbols", ", ".join(diagnostics.get("missing_earnings_symbols", [])) or "none"),
        ],
    )
    for caveat in diagnostics.get("data_caveats", []):
        doc.add_paragraph(caveat, style="List Bullet")

    doc.add_heading("Reproducibility", level=1)
    _add_kv_table(
        doc,
        [
            ("Run directory", str(run_dir)),
            ("Provider URI", str(manifest.get("provider_uri"))),
            ("Command", "; ".join(manifest.get("commands", []))),
        ],
    )

    out_path = out_path or (run_dir / "report.docx")
    doc.save(out_path)
    return out_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build DOCX reports for earnings factor run directories.")
    parser.add_argument("run_dirs", nargs="+", type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    for run_dir in args.run_dirs:
        out = build_report(run_dir)
        print(out)


if __name__ == "__main__":
    main()
