from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(31, 78, 121)
LIGHT = "D9EAF7"


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _fmt(value, pct: bool = False, digits: int = 2) -> str:
    if value is None:
        return "n/a"
    try:
        if pd.isna(value):
            return "n/a"
        value = float(value)
    except (TypeError, ValueError):
        return str(value)
    if pct:
        return f"{value:.{digits}%}"
    return f"{value:.{digits}f}"


def _set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _style_doc(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Calibri"
        styles[style_name].font.color.rgb = ACCENT
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)


def _short_name(name: str) -> str:
    mapping = {
        "EAP_v1_diversified_announcement_premium": "EAP",
        "PEAD_v1_quality_drift_60d": "PEAD60",
        "EBS_v1_leader_confirmed_basket_drift_10d": "EBS1_10",
        "EBS_v2_leader_follower_gap_40d": "EBS2_40",
    }
    for key, value in mapping.items():
        if name.startswith(key):
            suffix = name.removeprefix(key)
            return f"{value}{suffix}"
    return name


def _row(label: str, run_dir: Path) -> dict:
    manifest = _load_json(run_dir / "manifest.json")
    metrics = _load_json(run_dir / "metrics_summary.json")
    net = metrics.get("strategy_net", {})
    excess = metrics.get("excess_vs_QQQ", {})
    event = metrics.get("selected_event", {})
    ic = (
        metrics.get("signal_diagnostics", {})
        .get("selected", {})
        .get("rank_ic_rank_score_vs_strategy_excess_vs_QQQ_event_return")
    )
    return {
        "test": label,
        "factor": _short_name(manifest.get("factor_name", "")),
        "selected": metrics.get("event_counts", {}).get("selected_events"),
        "cum_ret": _fmt(net.get("cumulative_return"), pct=True),
        "ann_ret": _fmt(net.get("annualized_return"), pct=True),
        "sharpe": _fmt(net.get("sharpe", net.get("information_ratio"))),
        "max_dd": _fmt(net.get("max_drawdown"), pct=True),
        "excess_vs_QQQ": _fmt(excess.get("cumulative_return"), pct=True),
        "excess_IR": _fmt(excess.get("sharpe", excess.get("information_ratio"))),
        "avg_event": _fmt(event.get("average_event_return_after_cost"), pct=True),
        "avg_excess": _fmt(event.get("average_excess_vs_QQQ_event_return"), pct=True),
        "rank_IC": _fmt(ic, digits=3),
        "run_dir": str(run_dir),
    }


def _add_table(doc: Document, frame: pd.DataFrame, max_rows: int = 60) -> None:
    if frame.empty:
        doc.add_paragraph("No rows.")
        return
    frame = frame.head(max_rows)
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for idx, col in enumerate(frame.columns):
        table.rows[0].cells[idx].text = str(col)
        _set_cell_shading(table.rows[0].cells[idx], LIGHT)
        for paragraph in table.rows[0].cells[idx].paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(8)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[idx].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def build_report(cost_runs: list[tuple[str, Path]], pead_runs: list[tuple[str, Path]], out_path: Path) -> Path:
    doc = Document()
    _style_doc(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Earnings Factor Robustness Summary")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Cost sensitivity, leave-one-symbol, and leave-one-cluster checks").italic = True

    doc.add_heading("Decision", level=1)
    doc.add_paragraph(
        "PEAD60 remains the best candidate after robustness checks. It survives 50 bps one-way cost and "
        "the main leave-one-symbol tests, so it is not a single-name artifact like the retired pre-earnings "
        "read-through factor. Cluster tests show the signal remains positive without semiconductors or "
        "cybersecurity, while excluding enterprise software improves the result; this points to a need for "
        "cluster-specific calibration rather than one universal score."
    )

    doc.add_heading("Cost Sensitivity", level=1)
    cost = pd.DataFrame([_row(label, path) for label, path in cost_runs])
    cols = ["test", "factor", "selected", "cum_ret", "sharpe", "max_dd", "excess_vs_QQQ", "excess_IR", "avg_event"]
    _add_table(doc, cost[cols].sort_values(["factor", "test"]))

    doc.add_heading("PEAD60 Leave-One Robustness", level=1)
    pead = pd.DataFrame([_row(label, path) for label, path in pead_runs])
    cols = [
        "test",
        "selected",
        "cum_ret",
        "sharpe",
        "max_dd",
        "excess_vs_QQQ",
        "excess_IR",
        "avg_event",
        "avg_excess",
        "rank_IC",
    ]
    _add_table(doc, pead[cols], max_rows=30)

    doc.add_heading("Interpretation", level=1)
    for item in [
        "Cost: PEAD60 degrades but stays positive at 25 bps and 50 bps one-way cost.",
        "Single-name: no single tested top contributor fully explains the PEAD60 result.",
        "Cluster: enterprise software is currently a drag; semiconductors and cybersecurity both contribute useful but not exclusive alpha.",
        "Research status: PEAD60 is a research candidate, not production, until broader earnings coverage and PIT universe data are added.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build earnings factor robustness DOCX.")
    parser.add_argument("--cost", nargs="+", required=True, type=Path)
    parser.add_argument("--pead", nargs="+", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    cost_labels = [
        "base",
        "cost25",
        "cost50",
        "base",
        "cost25",
        "cost50",
        "base",
        "cost25",
        "cost50",
        "base",
        "cost25",
        "cost50",
    ]
    pead_labels = [
        "base",
        "no_crwd",
        "no_klac",
        "no_mu",
        "no_intc",
        "no_panw",
        "no_lrcx",
        "no_semiconductor",
        "no_semiconductor_equipment",
        "no_enterprise_software",
        "no_cybersecurity",
    ]
    if len(args.cost) != len(cost_labels):
        raise SystemExit(f"Expected {len(cost_labels)} cost runs, got {len(args.cost)}")
    if len(args.pead) != len(pead_labels):
        raise SystemExit(f"Expected {len(pead_labels)} PEAD runs, got {len(args.pead)}")
    print(build_report(list(zip(cost_labels, args.cost)), list(zip(pead_labels, args.pead)), args.out))


if __name__ == "__main__":
    main()
