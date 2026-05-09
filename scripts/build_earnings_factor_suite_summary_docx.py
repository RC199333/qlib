from __future__ import annotations

import argparse
import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(31, 78, 121)
LIGHT = "D9EAF7"


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _fmt(value, pct: bool = False, digits: int = 2) -> str:
    if value is None:
        return "n/a"
    try:
        if pd.isna(value):
            return "n/a"
        value = float(value)
    except (TypeError, ValueError):
        return str(value)
    if pct:
        return f"{value:.{digits}%}"
    return f"{value:.{digits}f}"


def _set_cell_shading(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _style_doc(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Calibri"
        styles[style_name].font.color.rgb = ACCENT
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.size = Pt(13)


def _add_table(doc: Document, frame: pd.DataFrame, max_rows: int = 40) -> None:
    if frame.empty:
        doc.add_paragraph("No rows.")
        return
    frame = frame.head(max_rows)
    table = doc.add_table(rows=1, cols=len(frame.columns))
    table.style = "Table Grid"
    for idx, col in enumerate(frame.columns):
        table.rows[0].cells[idx].text = str(col)
        _set_cell_shading(table.rows[0].cells[idx], LIGHT)
        for paragraph in table.rows[0].cells[idx].paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(8)
    for _, row in frame.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[idx].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def _short_name(name: str) -> str:
    mapping = {
        "EAP_v1_diversified_announcement_premium": "EAP",
        "PEAD_v1_quality_drift_60d": "PEAD60",
        "PEAD_v1_quality_drift_40d_h20": "PEAD20",
        "PEAD_v1_quality_drift_40d": "PEAD40",
        "PEAD_v1_quality_drift_40d_h60": "PEAD60_sweep",
        "EBS_v1_leader_confirmed_basket_drift_10d": "EBS1_10",
        "EBS_v1_leader_confirmed_basket_drift_h10": "EBS1_10_sweep",
        "EBS_v1_leader_confirmed_basket_drift": "EBS1_20",
        "EBS_v1_leader_confirmed_basket_drift_h40": "EBS1_40",
        "EBS_v2_leader_follower_gap_40d": "EBS2_40",
        "EBS_v2_leader_follower_gap_h10": "EBS2_10",
        "EBS_v2_leader_follower_gap": "EBS2_20",
        "EBS_v2_leader_follower_gap_h40": "EBS2_40_sweep",
    }
    return mapping.get(name, name)


def _row(run_dir: Path) -> dict:
    manifest = _load_json(run_dir / "manifest.json")
    metrics = _load_json(run_dir / "metrics_summary.json")
    spec = _load_json(run_dir / "experiment_spec.json")
    diagnostics = _load_json(run_dir / "data_diagnostics.json")
    net = metrics.get("strategy_net", {})
    excess = metrics.get("excess_vs_QQQ", {})
    event = metrics.get("selected_event", {})
    ic = (
        metrics.get("signal_diagnostics", {})
        .get("selected", {})
        .get("rank_ic_rank_score_vs_strategy_excess_vs_QQQ_event_return")
    )
    return {
        "label": _short_name(manifest.get("factor_name", "")),
        "factor": manifest.get("factor_name"),
        "hold_days": spec.get("execution", {}).get("hold_days"),
        "selected": metrics.get("event_counts", {}).get("selected_events"),
        "cum_ret": _fmt(net.get("cumulative_return"), pct=True),
        "ann_ret": _fmt(net.get("annualized_return"), pct=True),
        "sharpe": _fmt(net.get("sharpe", net.get("information_ratio"))),
        "max_dd": _fmt(net.get("max_drawdown"), pct=True),
        "excess_vs_QQQ": _fmt(excess.get("cumulative_return"), pct=True),
        "excess_IR": _fmt(excess.get("sharpe", excess.get("information_ratio"))),
        "avg_event": _fmt(event.get("average_event_return_after_cost"), pct=True),
        "avg_excess_event": _fmt(event.get("average_excess_vs_QQQ_event_return"), pct=True),
        "rank_IC": _fmt(ic, digits=3),
        "available_price_symbols": diagnostics.get("available_price_symbol_count"),
        "missing_earnings": len(diagnostics.get("missing_earnings_symbols", [])),
        "run_dir": str(run_dir),
    }


def build_summary(final_run_dirs: list[Path], sweep_run_dirs: list[Path], out_path: Path) -> Path:
    doc = Document()
    _style_doc(doc)
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    title = doc.add_paragraph()
    title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Earnings Factor Suite Summary")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Qlib-backed factor comparison and parameter decision").italic = True

    doc.add_heading("Decision", level=1)
    doc.add_paragraph(
        "The strongest current candidate is PEAD_v1_quality_drift_60d. EAP is a weak baseline, "
        "EBS_v1 is a weak but QQQ-relative positive basket-timing candidate, and EBS_v2 needs redesign "
        "because the leader-follower gap has weak rank IC and large drawdown."
    )

    doc.add_heading("Final Factor Comparison", level=1)
    final = pd.DataFrame([_row(path) for path in final_run_dirs])
    cols = [
        "label",
        "hold_days",
        "selected",
        "cum_ret",
        "ann_ret",
        "sharpe",
        "max_dd",
        "excess_vs_QQQ",
        "excess_IR",
        "avg_event",
        "avg_excess_event",
        "rank_IC",
    ]
    _add_table(doc, final[cols])

    doc.add_heading("Horizon Sweep", level=1)
    doc.add_paragraph(
        "The sweep was intentionally narrow: only economically meaningful holding windows were tested. "
        "This reduces the risk of turning the exercise into blind parameter mining."
    )
    sweep = pd.DataFrame([_row(path) for path in sweep_run_dirs])
    sweep_cols = ["factor", "hold_days", "selected", "cum_ret", "sharpe", "max_dd", "excess_vs_QQQ", "excess_IR"]
    sweep["factor"] = sweep["factor"].map(_short_name)
    _add_table(doc, sweep[sweep_cols].sort_values(["factor", "hold_days"]), max_rows=30)

    doc.add_heading("Data Limitation", level=1)
    doc.add_paragraph(
        "The local Qlib price universe contains more symbols than the current Alpha Vantage earnings cache. "
        "The API key hit the standard daily limit during expansion, so the final factor runs are complete "
        "for the currently cached earnings universe but not complete for the broader desired US tech pool."
    )
    coverage = final[["factor", "available_price_symbols", "missing_earnings"]].copy()
    coverage["factor"] = coverage["factor"].map(_short_name)
    _add_table(doc, coverage)

    doc.add_heading("Next Research Actions", level=1)
    for item in [
        "Backfill the missing earnings symbols when the API limit resets or switch to a higher-quality data source.",
        "Add revenue surprise, guidance revision, analyst revisions, and options-implied move.",
        "Replace current curated universe with point-in-time historical membership.",
        "Run leave-one-symbol and leave-one-cluster robustness before considering production.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Run Directories", level=1)
    run_dirs = final[["label", "run_dir"]].copy()
    _add_table(doc, run_dirs, max_rows=20)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the earnings factor suite comparison DOCX.")
    parser.add_argument("--final", nargs="+", required=True, type=Path)
    parser.add_argument("--sweep", nargs="+", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    print(build_summary(args.final, args.sweep, args.out))


if __name__ == "__main__":
    main()
